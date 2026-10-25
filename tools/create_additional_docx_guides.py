from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('deliverables')

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); b.append(e)
    tcPr.append(b)

def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None: node = OxmlElement(f'w:{side}'); mar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def keep(p):
    pPr = p._p.get_or_add_pPr(); keepNext = OxmlElement('w:keepNext'); pPr.append(keepNext)

def keep_lines(p):
    pPr = p._p.get_or_add_pPr(); keepLines = OxmlElement('w:keepLines'); pPr.append(keepLines)

def setup(title, subtitle):
    d = Document()
    s = d.sections[0]; s.top_margin = Inches(.78); s.bottom_margin = Inches(.72); s.left_margin = Inches(.82); s.right_margin = Inches(.82)
    styles = d.styles
    styles['Normal'].font.name = 'Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos'); styles['Normal'].font.size = Pt(10.5)
    styles['Normal'].paragraph_format.space_after = Pt(6); styles['Normal'].paragraph_format.line_spacing = 1.13
    for name, size in [('Title', 23), ('Subtitle', 12), ('Heading 1', 15), ('Heading 2', 11.5)]:
        st = styles[name]; st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'; st.font.color.rgb = RGBColor(0,0,0); st.font.size = Pt(size)
    styles['Heading 1'].font.bold = True; styles['Heading 1'].paragraph_format.space_before = Pt(15); styles['Heading 1'].paragraph_format.space_after = Pt(6)
    styles['Heading 2'].font.bold = True; styles['Heading 2'].paragraph_format.space_before = Pt(10); styles['Heading 2'].paragraph_format.space_after = Pt(4)
    p=d.add_paragraph(style='Title'); p.add_run(title); keep(p)
    p=d.add_paragraph(subtitle, style='Subtitle'); p.runs[0].font.color.rgb=RGBColor(80,80,80); p.paragraph_format.space_after=Pt(16)
    return d

def intro(d, heading, text):
    h=d.add_paragraph(heading, style='Heading 1'); keep(h); p=d.add_paragraph(text); keep_lines(p)

def section(d, heading, paragraphs):
    h=d.add_paragraph(heading, style='Heading 1'); keep(h)
    for text in paragraphs:
        p=d.add_paragraph(text); keep_lines(p)

def bullets(d, values):
    for v in values: d.add_paragraph(v, style='List Bullet')

def table(d, headers, rows, widths=None):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c, '1F4E79'); borders(c); set_cell_margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=val; borders(cells[i]); set_cell_margin(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n%2==1: shade(cells[i], 'F3F7FA')
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(9)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    d.add_paragraph().paragraph_format.space_after=Pt(2)

def footer(d):
    p=d.sections[0].footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('FocalDesk documentation guide  •  Repository grounded  •  25 October 2026'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)

def sources(d, names):
    section(d, 'Source documentation', ['This guide is a repository-grounded companion to the following maintained Markdown sources: ' + ', '.join(names) + '. Consult those files for implementation detail and the current command syntax.'])

def aios():
    d=setup('FocalDesk AIOS Platform Guide','Packages connectors context events routines and operational controls')
    intro(d,'Purpose and operating principle','This guide explains the FocalDesk AIOS platform for operators, package authors, and integrators. Its central rule is that packages, connectors, events, context, and suggestions carry bounded information or declarations; none grants authority by itself. Every execution path remains subject to explicit consent, capability, budget, and confirmation controls.')
    section(d,'Platform map',['AIOS is a collection of separate services and policy boundaries. Mission Control presents minimized operational state. Package Forge creates signed declarative packages; Package Manager inspects, trusts, stages, activates, and rolls them back. Event and Context Fabric constrain what evidence may enter the system, while Attention can turn consented events into inert suggestions.'])
    table(d,['Component','Purpose','Key boundary'],[
        ('Mission Control','Minimized timeline and runtime inventory','No raw tool arguments, results, context payloads, answers, or secrets'),
        ('Package Forge and Manager','Build and install signed FAI bundles','Signature validity is distinct from signer trust and activation'),
        ('Scenario Lab','Deterministic shadow evaluation','Never calls providers, runs tools, or mutates live services'),
        ('Event and Context Fabric','Consent based evidence intake and bounded context','Redaction, expiry, grants, and explicit field policies'),
        ('Attention and connectors','Suggestions and typed event production','Suggestions are inert; connectors cannot execute agents')], [1.55,2.15,3.1])
    section(d,'Package lifecycle',['A .fai bundle is declarative JSON. It may include agent, workflow, routine, connector, and Scenario Lab content, but cannot carry an executable, secret value, or instruction to silently grant authority. Inspect validates manifests, verifies the signature, evaluates included scenarios, and shows requested authority changes. Trusting a signer records a separate local decision. Staging requires a valid trusted signature, passing scenarios, and exact active dependencies; activation atomically selects the staged version and rebuilds registries.','After activation, every package-provided agent is disabled—even on update or rollback—until explicitly enabled again. Package activation also does not enable connectors, event sources, triggers, voice capture, or background execution. Use the Packages workspace or focaldesk-cli ai package commands to inspect before trust, stage, activate, or rollback.'])
    section(d,'Safe package authoring',['Generate protected signer material through Package Forge; the private seed remains in focald-secrets and only its opaque handle is used. Build processes test every included scenario and should refuse to overwrite outputs. In CI, run package tests on source projects and restrict signing to a protected release job; do not put a signing seed in a repository or ordinary CI variable.','A private registry is optional and loopback-only by default. It serves signed catalogs, pins catalog keys on clients, keeps coordinates immutable, and leaves downloaded packages in quarantine. A downloaded bundle must still be inspected, trusted, staged, and activated locally.'])
    section(d,'Evidence and consent',['Event Fabric accepts only enabled, typed sources with an explicit field allowlist, retention setting, and forwarding policy. It copies allowed scalar fields into a bounded redacted envelope; it does not scrape applications, pixels, clipboard, notification history, or calendar databases. Emergency disconnect rejects intake immediately without changing individual source policies.','Context Fabric contains selected typed desktop state or Console conversation, never screen pixels, keystrokes, secrets, arbitrary files, or clipboard contents. Publishing context does not share it with an agent. A user must make an expiring ContextGrant, which is intersected with the agent manifest at run start. Context is labeled untrusted evidence and cannot approve a mutation or expand a capability.'])
    section(d,'Connectors routines and recovery',['Connectors are bounded typed event producers. Installation and updates leave them disabled. Managed external connectors receive resource limits and transient user-service sandboxing; network access remains a separately granted permission. An authenticated connector event must still pass manifest, source-policy, field-disclosure, and emergency-disconnect checks.','Attention evaluates local, declarative routines and can emit an inert, expiring suggestion. A user must explicitly promote that suggestion into an agent or workflow run, which then passes normal capability lease, budget, permissions, and confirmation checks. Mission Control emergency pause suspends new triggers, routine dispatch, and Event Fabric intake; it does not cancel existing runs. Cancel runs by name and resume each affected function explicitly.'])
    sources(d,['docs/mission-control.md','docs/package-forge.md','docs/package-manager.md','docs/private-aios-registry.md','docs/scenario-lab.md','docs/context-fabric.md','docs/event-fabric.md','docs/attention-routines.md','docs/connector-host.md','docs/connector-sdk.md'])
    footer(d); d.save(OUT/'focaldesk-aios-platform-guide.docx')

def roadmap():
    d=setup('FocalDesk Product Roadmap','Current roadmap themes and remote desktop delivery direction')
    intro(d,'Purpose and status','This is a reader-friendly snapshot of the repository roadmap. It separates implemented alpha capabilities from planned work and should be read with the maintained Markdown roadmaps before commitments, release planning, or external statements are made.')
    section(d,'How to read this roadmap',['FocalDesk is alpha desktop software. Roadmap entries describe direction, dependencies, and safety boundaries rather than release promises. Hardware-dependent graphics, display, and capture work requires validation on representative GPUs, drivers, kernels, and display topologies.'])
    section(d,'Product themes',['The product direction combines a stable Wayland desktop foundation with optional services for workspaces, files, accessibility, display management, AI assistance, and remote access. The recurring design choice is to keep time-sensitive compositor work isolated from slower services, and to make sensitive automation and network-exposed functions opt-in.'])
    table(d,['Theme','Direction','Planning consideration'],[
        ('Desktop foundation','Refine compositor, shell, settings, services, and compatibility','Nested testing and recovery access remain the baseline'),
        ('Graphics and display','Advance HDR, color management, rendering, capture, and validation','Feature status varies by hardware and renderer'),
        ('AIOS','Evolve bounded agents, packages, context, events, and connectors','Consent, budgets, audit, and confirmation remain required'),
        ('Remote desktop','Expand from a constrained experimental service in deliberate phases','Network exposure and input authority require separate design gates')], [1.35,3.25,2.2])
    section(d,'Remote desktop baseline',['The current native remote-desktop service is experimental, view-only, loopback-only RDP for one output and one client. It has no input, clipboard synchronization, direct LAN exposure, independent logins, or production authentication. Use it only through an authenticated SSH tunnel with its short-lived owner-only token, and stop it when testing is complete.'])
    section(d,'Remote desktop delivery path',['The roadmap proceeds from the safe local baseline toward stronger transport and operational capabilities only after the preceding security and reliability gates are met. The intent is not to turn the current service into an unrestricted screen-sharing daemon. Work must preserve the opt-in model, explicit status, clear stop controls, bounded capture behavior, and recovery-oriented diagnostics.','The principal decisions to validate along the path are session ownership, authentication, authorization, encrypted transport, output and input policy, clipboard/file-transfer scope, auditing, multi-client behavior, and failure recovery. Each introduces a different user-trust and attack-surface question and should not be assumed from the presence of basic capture or RDP support.'])
    section(d,'Planning guidance',['Treat the roadmap as a sequencing aid. Before selecting a feature for a milestone, identify its safety boundary, test environment, rollback behavior, required service changes, and documentation impact. For remote desktop, include threat modeling and network review before enabling new exposure or input paths. For graphics work, include hardware matrix results and an SDR recovery path.'])
    sources(d,['ROADMAP.md','docs/remote-desktop-roadmap.md','docs/remote-desktop.md','docs/rendering.md','docs/hdr.md'])
    footer(d); d.save(OUT/'focaldesk-product-roadmap.docx')

def architecture():
    d=setup('FocalDesk Architecture and IPC Reference','System boundaries typed interfaces and replaceable shell clients')
    intro(d,'Purpose and central design','This reference is for engineers integrating with or changing FocalDesk. The compositor owns time-sensitive desktop state and rendering. Desktop services and shell clients are separate processes that communicate through typed IPC, so a slow service cannot take ownership of the frame loop.')
    section(d,'System responsibility boundaries',['The compositor handles Wayland protocol, surfaces, input routing, focus, window and workspace state, layout, outputs, damage, rendering, and presentation. Services own launch, settings, notifications, power, dialogs, controls, AI, automation, and other long-running workflows. Applications submit surfaces and protocol commits, not placement or scanout policy.'])
    table(d,['Layer','Owns','Does not own'],[
        ('Compositor engine','Display, input, workspace state, rendering, output coordination','Long-running policy workflows or service persistence'),
        ('Desktop services','Typed domain operations and durable service state','Compositor frame-loop authority'),
        ('Shell clients','User interfaces built over typed contracts','Direct mutation of compositor internals'),
        ('MCP and agents','Policy-constrained translation to approved IPC operations','An alternate unrestricted desktop API')], [1.35,3.05,2.4])
    section(d,'IPC contract',['FocalDesk IPC is private, typed, and capability-oriented. Request, response, and event types form the boundary between the compositor, services, and shell applications. A client should request a defined operation rather than reach into another component’s state. New interfaces should preserve validation, authorization, error handling, auditability, and version-conscious behavior at the boundary.','MCP sits on top of typed IPC as a policy and translation layer. Its catalog is intentionally read-heavy. Mutating calls require exact capability grants, explicit confirmation, and audit; it must never expose plaintext credentials.'])
    section(d,'Shell client model',['GTK shell clients are replaceable consumers of the system contracts. They should not duplicate compositor policy, invent a competing persistence model, or rely on private assumptions that make another shell impossible. Keep each client focused on its presentation and bounded interaction role, and locate shared domain decisions in the appropriate service or compositor contract.'])
    section(d,'Engineering implications',['When changing desktop behavior, first identify the owner: compositor for frame-critical state, a service for domain policy and persistence, or a shell client for presentation. Test typed boundaries directly where possible; then use nested compositor testing before a direct DRM session. Preserve recovery access and report hardware, backend, topology, and manual validation for graphics, capture, or HDR changes.'])
    sources(d,['docs/architecture.md','docs/ipc.md','docs/architecture/trusted-shell-clients.md','docs/engineering.md','docs/mcp.md','README.md'])
    footer(d); d.save(OUT/'focaldesk-architecture-and-ipc-reference.docx')

def security():
    d=setup('FocalDesk Security and Privacy Guide','Trust boundaries credentials consent automation and reporting')
    intro(d,'Purpose and central rule','This guide explains how FocalDesk limits authority and sensitive data. Use Settings and the typed service interfaces for ordinary work. Treat AI, automation, remote access, connectors, and credential operations as opt-in capabilities with separate consent, confirmation, and audit boundaries.')
    section(d,'Security model',['FocalDesk separates time-sensitive compositor work from service policy, uses private typed IPC for trusted components, and defaults sensitive features to disabled or narrowly scoped. A feature being installed does not mean it is authorized, network-enabled, or active. Keep recovery access to a known-good desktop or TTY when experimenting with alpha session software.'])
    table(d,['Area','Protection','Operator expectation'],[
        ('Credentials','Brokered opaque handles and ACL-protected access','Never place keys in settings files, screenshots, shell history, or tickets'),
        ('Agents and MCP','Capabilities, budgets, confirmation, and audit','Review manifests and tools; grant only required scope'),
        ('Events and context','Consent, minimization, expiry, and grants','Enable only necessary sources and fields; clear data when no longer needed'),
        ('Connectors and remote access','Disabled by default and separately controlled','Avoid broad network exposure; retain explicit stop controls')], [1.35,2.75,2.7])
    section(d,'Credentials and secret handling',['focald-secrets supplies a native ACL-protected socket and the standard Secret Service API. Secrets remain opaque to the AI and MCP layers. Provider credentials are stored at designated secret handles; environment variables are development or upgrade fallbacks, not a reason to put credentials into general configuration.','If migrating from GNOME Keyring, preview the copy operation first. Both implementations cannot own org.freedesktop.secrets simultaneously. The migration copies readable entries rather than deleting originals; review secrets and service ownership before changing the active broker.'])
    section(d,'Consent privacy and automation',['Event Fabric copies only explicitly permitted scalar fields into bounded redacted envelopes. Context Fabric does not collect screen pixels, clipboard contents, keystrokes, arbitrary files, or secret values. Agents require an expiring context grant and still cannot treat context as authority. Attention routines can create suggestions but cannot start a run or mutation.','Automation, connector, and remote-desktop functions are intentionally separate opt-in systems. An agent or package cannot grant itself authority. For mutations, preserve native one-shot confirmation and audit behavior rather than substituting a broad remembered approval.'])
    section(d,'Reporting vulnerabilities',['Report potential vulnerabilities privately through the process in SECURITY.md. Include a clear description, affected version or commit, reproduction details, impact, and any suggested mitigation. Do not publish exploit details before maintainers have had an opportunity to investigate and coordinate a fix.'])
    section(d,'Operational checks',['Use focaldesk-cli diagnostics --no-logs for a bounded diagnostic bundle without log contents. Review bundles before sharing because metadata can still be sensitive. Check service health with systemctl --user --failed and examine focused logs with journalctl; avoid copying credentials, tokens, or personal data into issue reports.'])
    sources(d,['SECURITY.md','docs/secrets.md','docs/mcp.md','docs/event-fabric.md','docs/context-fabric.md','docs/connector-sdk.md','docs/remote-desktop.md'])
    footer(d); d.save(OUT/'focaldesk-security-and-privacy-guide.docx')

if __name__ == '__main__':
    OUT.mkdir(exist_ok=True)
    aios(); roadmap(); architecture(); security()
