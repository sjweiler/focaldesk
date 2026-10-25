#!/usr/bin/env python3
"""Create the remaining FocalDesk end-user and operator guides."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'deliverables'
BLUE='17365D'; PALE='F3F7FB'

GUIDES={
'focaldesk-daily-user-guide.docx':('FocalDesk Daily User Guide','Workspaces shortcuts settings files and recovery',[
('Getting started',['FocalDesk is an alpha Wayland desktop. Use Settings to make ordinary changes and keep a working alternate desktop session while becoming familiar with it. The Launcher, workspace controls, dock, panel, file manager, and Settings application are intended to work together rather than as separate configuration systems.']),
('Essential shortcuts',['Super+Enter opens the configured terminal; Super+B opens the browser; Super+Q closes the focused window; Super+L locks the session; Alt+Tab cycles applications; Ctrl+Alt+D toggles the Launcher. Super++ and Super+- adjust the focused display magnifier.']),
('Workspace workflow',['Alt+1 through Alt+9 switch to numbered workspaces. Alt+Shift+1 through Alt+Shift+9 move the focused window and follow it. Alt+0 opens the complete workspace list. Each output can show a different workspace. Enable Restore session in Settings → Workspaces only if you want compatible applications and window placement restored at login.']),
('Window layout',['Split screen is disabled by default. Enable it in Settings → Workspaces. Super+Left and Super+Right create a side-by-side split; Super+Up and Super+Down create stacked panes; Super+Z opens the layout chooser. Super+Alt+Backspace undoes the last split action.']),
('Everyday settings',['Use Settings for appearance, display, input, applications, workspaces, privacy, power, accessibility, and service controls. The canonical configuration is $XDG_CONFIG_HOME/focaldesk/settings.json. Avoid hand-editing it unless necessary; preserve unknown keys, make a recovery copy, validate JSON, and reload only after a valid save.']),
('Files notifications and privacy',['FocalDesk Files supports local browsing, mounts, network locations, archive extraction, undo, trash, and bookmarks. Privacy controls include recent activity and AI permission records. Clear recent activity from Settings if you do not want previous application launches and file opens shown in Launcher search.']),
('When something goes wrong',['Use a TTY or a known-good desktop session if the display is unusable. Review user services with systemctl --user --failed. Use focaldesk-cli diagnostics --no-logs for a bounded diagnostic bundle without log content. Do not change generated displays.json while FocalDesk is running.'])]),
'focaldesk-administrator-operations-guide.docx':('FocalDesk Administrator Operations Guide','Services diagnostics upgrades backups and recovery',[
('Operational model',['FocalDesk is installed as a Wayland session plus system and user-session components. Fedora’s system-installed user units live under /usr/lib/systemd/user but are managed with systemctl --user because they need the graphical user bus and runtime directory.']),
('Routine health checks',['Check the session target and failed user units after login. Use journalctl -b -t focaldesk for compositor events and journalctl --user -u SERVICE -b for a specific service. Review diagnostic bundles before sharing because they can contain private metadata.']),
('Upgrade discipline',['Use the committed Cargo.lock and repository installation recipes. Do not use cargo update as a general repair operation. For Fedora system layout, avoid mixing old ~/.config/systemd/user development units with packaged units; old user units can override system units.']),
('Backup scope',['Back up $XDG_CONFIG_HOME/focaldesk, $XDG_DATA_HOME/focaldesk, $XDG_STATE_HOME/focaldesk, and $XDG_CACHE_HOME/focaldesk when preserving an alpha installation. Treat credential files as sensitive and do not copy them into tickets or public archives.']),
('Display manager recovery',['Keep GDM installed when testing focaldmd. If focaldmd fails, switch to a TTY, disable focaldmd.service, enable gdm.service, and confirm display-manager.service. Do not delete GDM or edit its packaged configuration to install FocalDesk.']),
('Removal',['The repository uninstaller removes paths listed in reviewed manifests and preserves user state by default. The purge option deletes FocalDesk user configuration and state and is irreversible; create a backup first.'])]),
'focaldesk-ai-and-credentials-guide.docx':('FocalDesk AI and Credentials Guide','Providers permissions credential broker and agent safety',[
('Security boundary',['AI, automation, agents, and credentials are permissioned experimental features. The broker never exposes plaintext secrets through MCP. Use the permission UI for normal decisions and do not put keys in settings files, screenshots, shell history, or shared logs.']),
('Credential broker',['focald-secrets provides a native, ACL-protected FocalDesk socket and the standard org.freedesktop.secrets API. On Fedora with focaldmd, PAM unlocks the broker at login. Its packaged ACL is default-deny for native service access.']),
('Provider setup',['Cloud-backed providers look for ai/openai-api-key, ai/anthropic-api-key, or ai/vllm-api-key in focald-secrets. Environment variables remain development and upgrade fallbacks. Restart focaldesk-server.service after changing a stored provider key.']),
('GNOME Keyring migration',['GNOME Keyring and focald-secrets cannot own org.freedesktop.secrets simultaneously. While GNOME Keyring is active, preview then optionally import readable entries with focald-secrets-import-gnome-keyring --dry-run and --import. The import copies rather than deletes originals.']),
('Agents and permissions',['Agent manifests are declarative and live below $XDG_CONFIG_HOME/focaldesk/agents. They can narrow audited tools and budgets but cannot grant themselves authority or bypass one-shot confirmation. Enable only agents whose instructions, tools, scopes, and triggers you understand.']),
('Verification',['Check the credential broker with systemctl status focald-secrets@UID.service, inspect the Secret Service owner with busctl --user status org.freedesktop.secrets, and use scripts/check-secrets-health.sh from the source checkout.'])]),
'focaldesk-display-and-graphics-guide.docx':('FocalDesk Display and Graphics Guide','Displays scaling HDR renderers and graphics recovery',[
('Choose a safe baseline',['Start with one SDR display at 60 Hz on the GLES renderer. Add native resolution, hardware cursor, additional outputs, high refresh, 10-bit scanout, and HDR one change at a time. Keep a working alternate session and a TTY available.']),
('Renderer choice',['GLES is the packaged compatibility and recovery default. Raw Ash Vulkan is the direct DRM/KMS renderer for validated hardware. The nested wgpu Vulkan backend is a development backend, not the production DRM path. Change focaldmd renderer selection only in the existing /etc/focaldmd.toml document.']),
('Display settings',['Use Settings for connector mode, position, scale, primary output, ICC profile, HDR request, and per-output appearance. displays.json is generated output topology; do not edit it while the compositor runs.']),
('HDR expectations',['HDR and color management are experimental and hardware-dependent. Test HDR only on validated outputs. Preserve SDR usability, verify color appearance after every change, and do not enable undocumented HDR or atomic-commit overrides as a generic troubleshooting step.']),
('Capture and screenshots',['PipeWire/portal capture is experimental and relies on the portal setup. Direct DRM screenshots are 16-bit Display P3 PNGs with an ICC profile. Capture paths may intentionally differ from HDR output encoding to retain a usable SDR interchange format.']),
('Recovery evidence',['For a blank display record GPU, driver, kernel, connector names, modes, scales, transforms, HDR state, and the last journal messages. Use journalctl -b -t focaldesk.'])]),
'focaldesk-developer-guide.docx':('FocalDesk Developer Guide','Build nested testing services IPC and contribution checks',[
('Development approach',['FocalDesk is a Rust workspace. Build on Fedora, test nested before attempting a direct DRM session, and retain a recovery desktop. The nested winit backend is the established compatibility test path; it avoids risking the current graphical session.']),
('Build and test',['Install the Fedora build prerequisites from docs/building.md, install Rust with rustup, then run cargo build --workspace. The baseline checks are cargo fmt --all -- --check, cargo check --workspace, cargo clippy --workspace --all-targets, and cargo test --workspace.']),
('Nested execution',['Run cargo run -p focaldesk-desktop --no-default-features --features winit,xwayland for a nested session. Run just nested-smoke for bounded smoke artifacts. Use just nested-wgpu only when working on the experimental nested Vulkan backend.']),
('Code boundaries',['The compositor owns display, input, windows, workspaces, and rendering. Services own launch, settings, notifications, power, dialogs, controls, AI, and automation through typed IPC. Preserve this boundary so slow services cannot stall the frame loop.']),
('MCP and agents',['focaldesk-mcp is a policy and translation layer over typed IPC, not an alternate desktop API. Its catalog is read-heavy; mutation requires exact capability grants, explicit confirmation, and audit. Never add a tool that returns plaintext credentials.']),
('Change validation',['Use targeted unit tests, architecture checks, nested smoke tests, and hardware regression where applicable. For compositor, capture, or HDR changes report GPU, driver, backend, topology, and manual validation performed.'])]),
'focaldesk-automation-and-remote-desktop-guide.docx':('FocalDesk Automation and Remote Desktop Guide','Opt in controls authority boundaries and safe operation',[
('Why these features are separate',['Automation and remote desktop can cause meaningful actions or expose the desktop. They are intentionally opt-in and should not be enabled merely because their binaries are installed. Review configuration and scripts before activation.']),
('Automation',['The command-executing automation service is not included in the default service bundle. Install it explicitly only after reviewing automations.toml and referenced Lua scripts. Automation definitions live below $XDG_CONFIG_HOME/focaldesk/automation and are security-sensitive.']),
('Agent triggers',['Declarative agent triggers may run on schedules or desktop events but cannot bypass the normal permission and one-shot mutation-confirmation path. Limit tool allowlists, scopes, budgets, cooldowns, and hourly run counts. Suspend triggers when investigating unexpected behavior.']),
('Remote desktop boundary',['focaldesk-remoted is experimental, view-only, loopback-only RDP for one output and one client. Input, clipboard sync, direct LAN exposure, independent logins, and production authentication are not implemented.']),
('Safe remote use',['Start the service explicitly, read the short-lived token from its owner-only runtime file, and connect through an authenticated SSH tunnel. Never forward port 3389 from a router or change the listener to all interfaces. Stop the service when testing is complete.']),
('Logs and stop controls',['Stop remote capture with systemctl --user stop focaldesk-remoted.service and inspect journalctl --user -u focaldesk-remoted.service -b. Stop automation by disabling its user service. Use the microphone kill control for FocalDesk microphone clients; it is not a hardware mute for unrelated applications.'])])}

def shade(cell,color):
    p=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd');e.set(qn('w:fill'),color);p.append(e)
def table(doc,rows):
    t=doc.add_table(rows=1,cols=2);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(('Topic','Guidance')):
        c=t.rows[0].cells[i];c.width=Inches(1.8 if i==0 else 5.5);shade(c,BLUE);c.text=h;c.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255);c.paragraphs[0].runs[0].bold=True
    for n,(a,b) in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate((a,b)):
            cells[i].width=Inches(1.8 if i==0 else 5.5);shade(cells[i],PALE if n%2 else 'FFFFFF');cells[i].text=v;cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for r in cells[i].paragraphs[0].runs:r.font.size=Pt(9.5)
    doc.add_paragraph()
def make(path,title,sub,sections):
    d=Document();s=d.sections[0];s.top_margin=Inches(.75);s.bottom_margin=Inches(.75);s.left_margin=Inches(.78);s.right_margin=Inches(.78)
    st=d.styles;st['Normal'].font.name='Aptos';st['Normal'].font.size=Pt(10.5);st['Title'].font.name='Aptos Display';st['Title'].font.size=Pt(24);st['Title'].font.color.rgb=RGBColor(0,0,0)
    for b in st['Title']._element.get_or_add_pPr().findall(qn('w:pBdr')):st['Title']._element.get_or_add_pPr().remove(b)
    for x in ('Heading 1','Heading 2'):st[x].font.name='Aptos';st[x].font.color.rgb=RGBColor(0,0,0)
    d.add_paragraph(title,style='Title');p=d.add_paragraph(sub);p.runs[0].font.size=Pt(12);p.runs[0].font.color.rgb=RGBColor(70,70,70);p.paragraph_format.space_after=Pt(14)
    d.add_heading('Purpose',1);p=d.add_paragraph(f'This guide is a concise, repository-grounded reference for {title.removeprefix("FocalDesk ").removesuffix(" Guide").lower()}. FocalDesk remains alpha software; retain recovery access and use the project’s typed Settings and service interfaces for ordinary work.');p.paragraph_format.space_after=Pt(9)
    table(d,sections)
    d.add_heading('Useful commands',1);p=d.add_paragraph('Run these read-only checks from the active user session when relevant:');p.paragraph_format.space_after=Pt(3)
    c=d.add_paragraph('focaldesk-cli diagnostics --no-logs\nsystemctl --user --failed\njournalctl -b -t focaldesk',style='Normal')
    for r in c.runs:r.font.name='Consolas';r.font.size=Pt(9)
    d.add_heading('Sources',1);d.add_paragraph('Repository sources reviewed: README.md, docs/building.md, docs/configuration.md, docs/troubleshooting.md, and the topic-specific documentation in docs/. Repository state reviewed on 2026 10 24.')
    h=s.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;h.add_run(title).font.size=Pt(8)
    f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.add_run(title).font.size=Pt(8)
    d.save(path)
def main():
    OUT.mkdir(exist_ok=True)
    for f,(t,sub,sec) in GUIDES.items():make(OUT/f,t,sub,sec);print(OUT/f)
if __name__=='__main__':main()
