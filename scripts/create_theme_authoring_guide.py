#!/usr/bin/env python3
"""Create a concise FocalDesk theme authoring guide."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables' / 'focaldesk-theme-authoring-guide.docx'

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tcPr.append(shd)

def borders(table):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        x = OxmlElement(f'w:{edge}'); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); b.append(x)
    tblPr.append(b)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr(); hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'),'true'); trPr.append(hdr)

def cell_text(cell, text, bold=False, white=False):
    cell.text = ''; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); r = p.add_run(text); r.bold = bold; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(255,255,255) if white else RGBColor(0,0,0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; borders(t)
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.width=Inches(widths[i]); shade(c,'17365D'); cell_text(c,h,True,True)
    set_repeat_table_header(t.rows[0])
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].width=Inches(widths[i]); shade(cells[i], 'F3F7FB' if n%2 else 'FFFFFF'); cell_text(cells[i],v)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def code(doc, text):
    p=doc.add_paragraph(); p.style='Code'; p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(8); p.paragraph_format.keep_together=True
    for i,line in enumerate(text.split('\n')):
        r=p.add_run(line); r.font.name='Consolas'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Consolas'); r.font.size=Pt(9)
        if i < len(text.split('\n'))-1: r.add_break()

def heading(doc, text, level=1): doc.add_heading(text, level)
def para(doc, text):
    p=doc.add_paragraph(text); p.paragraph_format.space_after=Pt(7); return p

def bullet(doc, text):
    p=doc.add_paragraph(style='List Bullet'); p.add_run(text); p.paragraph_format.space_after=Pt(3)

def main():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.72); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5)
    styles['Title'].font.name='Aptos Display'; styles['Title'].font.size=Pt(24); styles['Title'].font.color.rgb=RGBColor(0,0,0)
    title_ppr = styles['Title']._element.get_or_add_pPr()
    for border in title_ppr.findall(qn('w:pBdr')): title_ppr.remove(border)
    for name,size in [('Heading 1',16),('Heading 2',12)]: styles[name].font.name='Aptos'; styles[name].font.size=Pt(size); styles[name].font.bold=True; styles[name].font.color.rgb=RGBColor(0,0,0)
    code_style=styles.add_style('Code',1); code_style.font.name='Consolas'; code_style.font.size=Pt(9); code_style.paragraph_format.left_indent=Inches(.12); code_style.paragraph_format.right_indent=Inches(.12)
    title=doc.add_paragraph(style='Title'); title.add_run('FocalDesk Theme Authoring Guide')
    sub=doc.add_paragraph('Theme Editor workflow documents packages wallpaper and safe testing'); sub.runs[0].font.size=Pt(12); sub.runs[0].font.color.rgb=RGBColor(64,64,64); sub.paragraph_format.space_after=Pt(18)
    heading(doc,'Purpose')
    para(doc,'This guide explains how to create, preview, save, package, import, and remove FocalDesk themes. The recommended workflow is to author in Settings → Theme Editor, preserve a source TOML file, preview changes temporarily, then export a portable .fdtheme package for installation and sharing. Theme formats are versioned alpha interfaces, so retain source files across upgrades.')
    heading(doc,'What a theme controls')
    table(doc,['Area','Examples','Authoring goal'],[
        ['Paint and wallpaper','Solid, linear, or radial paint; fit, dim, tint, blur, saturation','Set the desktop base and retain legibility behind shell surfaces.'],
        ['Semantic surfaces','Bar, dock, button, active button, popup, window frame','Use role-based colors rather than one flat palette.'],
        ['Interaction states','Hover, pressed, selected, focused, urgent, disabled','Keep feedback visible in every state.'],
        ['Edges and type','Border, highlight, shadow, glow, radius, font size and weight','Create hierarchy without sacrificing text contrast.'],
        ['Layout and color behavior','Bar height, dock width, spacing, icon size, gamut and dynamic range','Keep targets usable and choose SDR or HDR intentionally.'],
    ],[1.35,2.5,3.9])
    heading(doc,'Start in Theme Editor')
    para(doc,'Open Settings → Theme Editor. The editor is organized into General, Surfaces, Layout and Color, Wallpaper, and Paint pages. It can author sRGB and Display P3 colors, SDR or HDR intent, gradients with two through eight ordered stops, wallpaper processing, typography, and shell metrics. The state-matrix preview and compositor contrast status identify primary-text issues, but they do not silently change your colors.')
    bullet(doc,'Begin from Default or import an existing package as a reference; keep the source outside the installed package directory.')
    bullet(doc,'Change the semantic role that needs adjustment. For example, change active_button for selected controls instead of making every surface brighter.')
    bullet(doc,'Use Live preview for rapid evaluation. It is temporary; Apply affects the current compositor process only, while Save writes the TOML source.')
    heading(doc,'Save preview and apply')
    table(doc,['Action','What it changes','Persistence'],[
        ['Save','Writes the ThemeDocument TOML source','Persistent source file; does not change the running desktop by itself.'],
        ['Live preview','Debounced valid edits sent to the compositor','Temporary; removed when leaving the editor.'],
        ['Apply','Sets the running compositor theme','Runtime only; not retained through compositor restart.'],
        ['Revert Preview','Restores the last applied runtime theme','Temporary preview is discarded.'],
        ['Export','Creates validated portable .fdtheme package','Portable artifact containing theme document and optional wallpaper.'],
    ],[1.15,3.0,3.6])
    doc.add_page_break()
    heading(doc,'Theme document structure',1)
    para(doc,'The current TOML document format is version 1. A standalone document references its wallpaper by filesystem path; moving the document or the wallpaper can break that reference. Use a package when the theme needs to travel with its wallpaper.')
    code(doc,'''format_version = 1
name = "My Theme"
dynamic_range = "sdr"
hdr_luminance_nits = 1000.0

[paint]
mode = "solid"

[paint.color]
space = "srgb"
r = 0.04
g = 0.20
b = 0.73
a = 1.0

[wallpaper]
path = "/absolute/path/to/wallpaper.png"
fit = "fill"
dim = 0.0''')
    heading(doc,'Semantic tokens and state design')
    para(doc,'Use the semantic token groups to make the desktop consistent. The shipped Default theme includes surfaces for bar, dock, button, active_button, popup, and window_frame. Each surface can have normal, hover, pressed, selected, focused, urgent, or disabled behavior as supported by the editor and compositor. Inherited states prevent needless duplication, but deliberate state colors are important for focus, danger, and unavailable controls.')
    table(doc,['Token group','Use it for','Practical check'],[
        ['semantic.surfaces','Shell background roles','Can text and icons be read over both wallpaper and surface?'],
        ['semantic.edges','Borders, inner highlight, shadow, glow, radius','Are bounds visible without looking like hard boxes everywhere?'],
        ['semantic.typography','Primary and secondary text, size, weight, letter spacing','Is secondary text subdued but still readable?'],
        ['semantic.layout','Bar, dock, padding, gap, icon size','Are touch and pointer targets usable at the selected scale?'],
        ['semantic.color_behavior','Gamut, dynamic range, white and luminance limits, mapping','Does the intent match the display and desktop use case?'],
    ],[1.55,2.25,3.95])
    heading(doc,'Wallpaper and HDR considerations')
    para(doc,'A package can embed at most one wallpaper. PNG, JPEG, and WebP are treated as SDR sRGB. OpenEXR is imported as linear HDR with Rec.709 or sRGB primaries in the GLES renderer. A 16-bit PNG tagged Rec.2020 PQ or HLG through cICP metadata is decoded to the same linear HDR scene. JPEG is SDR unless a reproducible Ultra HDR gain-map decoder is included. The editor refuses unsafe wallpaper references and invalid theme values before preview.')
    table(doc,['Constraint','Limit or behavior'],[
        ['Supported packaged wallpaper formats','PNG, JPEG, WebP, OpenEXR'],
        ['Maximum asset size','64 MiB'],
        ['Maximum dimensions','32,768 pixels on either axis; 100 million pixels total'],
        ['Package verification','Records asset size and SHA 256 digest; validates both before installation'],
        ['HDR authoring','Use only on a tested HDR output; preserve SDR usability and contrast'],
    ],[2.2,5.55])
    heading(doc,'Export import and select',1)
    para(doc,'Use Theme Editor Export to create an .fdtheme package. Import validates the package and installs it below $XDG_DATA_HOME/focaldesk/themes/<theme-slug>/, normally ~/.local/share/focaldesk/themes/. The installed directory contains theme.toml and the packaged wallpaper. The installer will not follow existing symlinks at either destination.')
    code(doc,'''# Typical installed-theme location
theme_root="${XDG_DATA_HOME:-$HOME/.local/share}/focaldesk/themes"
find "$theme_root" -maxdepth 2 -type f -print

# Choose an imported theme in Settings → Appearance → Theme.
# Use Theme Editor → Uninstall to remove a package.''')
    para(doc,'The Default theme is the system document /usr/share/focaldesk/default.toml and loads when appearance.theme is absent or set to Default. Optional source packages are available under assets/themes/optional/; import them through Theme Editor rather than copying their contents into system directories.')
    heading(doc,'Safe authoring checklist')
    bullet(doc,'Keep the source TOML and wallpaper in a version-controlled or backed-up location before export.')
    bullet(doc,'Preview every interaction state and check primary text in bar, dock, popup, and buttons.')
    bullet(doc,'Test on the intended display scale and, for HDR themes, with SDR content and the target HDR output.')
    bullet(doc,'Export before installing; import the package rather than hand-copying files into the XDG themes directory.')
    bullet(doc,'If the preview is undesirable, use Revert Preview or leave Theme Editor. If an installed theme fails to load, FocalDesk selects Default safely.')
    heading(doc,'Troubleshooting and recovery')
    para(doc,'A valid editor preview is not a persistent save. If a theme disappears after restart, confirm it was exported, imported, and selected in Settings → Appearance. If a standalone TOML loses its wallpaper, repair the path in the source document or export it as a package. Do not hand-edit an installed package while FocalDesk is running. For a broken appearance configuration, select Default in Settings or remove the package through Theme Editor; keep your original source TOML as the recovery copy.')
    heading(doc,'Repository sources')
    para(doc,'This guide is based on docs/theme-editor.md, docs/configuration.md, assets/themes/default.toml, assets/themes/focaldesk-theme.toml, and the optional .fdtheme packages in assets/themes/optional/. Repository state reviewed on 2026 10 24.')
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=header.add_run('FocalDesk Theme Authoring Guide'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=footer.add_run('FocalDesk Theme Authoring Guide'); r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)
    OUT.parent.mkdir(exist_ok=True); doc.save(OUT); print(OUT)

if __name__=='__main__': main()
